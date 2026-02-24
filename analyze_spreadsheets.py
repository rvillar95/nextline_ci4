import os
import sys
from openpyxl import load_workbook


FILES = [
    r"c:\wamp64\www\codeigniter4\nextline_ci4\Planilla Cálculo segun porciones.xlsx",
    r"c:\wamp64\www\codeigniter4\nextline_ci4\CALCULO DE PORCIONES.xlsx",
    r"c:\wamp64\www\codeigniter4\nextline_ci4\FICHA INGRESO ADULTO 2026 (1).xlsm",
]


def cell_str(v) -> str:
    if v is None:
        return ""
    s = str(v).replace("\n", " ").strip()
    return s[:140]


def summarize_sheet(ws):
    max_r = ws.max_row or 0
    max_c = ws.max_column or 0

    # collect non-empty cells in a top-left window
    non_empty = []
    for r in range(1, min(80, max_r) + 1):
        for c in range(1, min(50, max_c) + 1):
            v = ws.cell(r, c).value
            if v not in (None, ""):
                non_empty.append((r, c, cell_str(v)))

    # header candidates: rows with >=4 non-empty cells in first 25 columns
    header_rows = []
    for r in range(1, min(60, max_r) + 1):
        cnt = 0
        for c in range(1, min(25, max_c) + 1):
            v = ws.cell(r, c).value
            if v not in (None, ""):
                cnt += 1
        if cnt >= 4:
            header_rows.append((r, cnt))

    # sample formulas
    formulas = []
    for r in range(1, min(300, max_r) + 1):
        for c in range(1, min(80, max_c) + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.startswith("="):
                formulas.append((r, c, cell_str(v)))
                if len(formulas) >= 40:
                    break
        if len(formulas) >= 40:
            break

    return {
        "title": ws.title,
        "max_row": max_r,
        "max_col": max_c,
        "header_rows": header_rows[:12],
        "non_empty_top": non_empty[:120],
        "formula_samples": formulas[:20],
    }


def main():
    # Avoid Windows console encoding issues with accents/combining marks
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="backslashreplace")
    except Exception:
        pass

    for path in FILES:
        print("\n" + "=" * 100)
        print("FILE:", path)
        if not os.path.exists(path):
            print("  !! NOT FOUND")
            continue

        keep_vba = path.lower().endswith(".xlsm")
        wb = load_workbook(path, data_only=False, keep_vba=keep_vba)
        print("Sheets:", wb.sheetnames)

        # defined names (helpful to understand key calc inputs/outputs)
        try:
            names = sorted(list(wb.defined_names))
            print("Defined names:", len(names))
            for n in names[:60]:
                print(" -", n)

            # If there is a named range 'porciones', dump its bounds and sample rows
            if "porciones" in wb.defined_names:
                dn = wb.defined_names["porciones"]
                print("\nNamed range 'porciones' destinations:")
                for title, coord in dn.destinations:
                    print(f" - {title}!{coord}")
                    ws = wb[title]
                    cells = ws[coord]
                    # cells is 2D tuple
                    max_r = min(len(cells), 15)
                    max_c = min(len(cells[0]) if cells else 0, 10)
                    print("   Sample (first rows):")
                    for r in range(max_r):
                        row = [cell_str(cells[r][c].value) for c in range(max_c)]
                        if any(row):
                            print("   ", row)
        except Exception as e:
            print("Defined names: error", e)

        for sname in wb.sheetnames[:15]:
            ws = wb[sname]
            info = summarize_sheet(ws)
            print("\n-- SHEET:", info["title"], f"(rows={info['max_row']}, cols={info['max_col']})")
            if info["header_rows"]:
                print("Header-like rows:", info["header_rows"])
            print("Top non-empty cells sample (r,c,val):")
            for r, c, v in info["non_empty_top"][:30]:
                print(f"  ({r},{c}) {v}")
            if info["formula_samples"]:
                print("Formula samples:")
                for r, c, v in info["formula_samples"][:12]:
                    print(f"  ({r},{c}) {v}")


if __name__ == "__main__":
    main()

